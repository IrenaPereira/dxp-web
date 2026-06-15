#!/usr/bin/env python3
"""Generate the Digital Experiments capabilities deck as a native .pptx.

Mirrors deck/index.html 1:1 (content) in the current brand frame.
Run:  python3 deck/build_pptx.py
Out:  deck/Digital-Experiments-Capabilities.pptx
"""
import os
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

# ---- Brand tokens -----------------------------------------------------------
INK      = RGBColor(0x14, 0x13, 0x1A)
INK2     = RGBColor(0x1D, 0x1D, 0x1D)
PAPER    = RGBColor(0xFF, 0xFF, 0xFF)
PAPER2   = RGBColor(0xF6, 0xF5, 0xF2)
LINE     = RGBColor(0xE3, 0xE1, 0xE4)
MUTED    = RGBColor(0x6C, 0x6A, 0x76)
MUTED2   = RGBColor(0x9B, 0x99, 0xA4)
ACCENT   = RGBColor(0x5B, 0x53, 0xE0)
ACCENT3  = RGBColor(0xB9, 0xB5, 0xF4)
ONINK    = RGBColor(0xF5, 0xF4, 0xF7)
ONINKMUT = RGBColor(0xB5, 0xB2, 0xC2)
ONINKLN  = RGBColor(0x32, 0x2F, 0x3D)
CARDINK  = RGBColor(0x20, 0x1F, 0x29)

DISPLAY = "Playfair Display"
SANS    = "Inter"
MONO    = "JetBrains Mono"

EMU_IN = 914400
SW, SH = 13.333, 7.5
ML = 0.92                      # left/right margin
CW = SW - 2 * ML               # content width

prs = Presentation()
prs.slide_width  = Emu(int(SW * EMU_IN))
prs.slide_height = Emu(int(SH * EMU_IN))
BLANK = prs.slide_layouts[6]


def slide(bg):
    s = prs.slides.add_slide(BLANK)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = bg
    return s


def _set_space_after(p, pts):
    p.space_after = Pt(pts)


def tx(slide, l, t, w, h, anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    return tf


def para(tf, runs, size, color=None, bold=False, font=SANS, align=PP_ALIGN.LEFT,
         space_after=6, space_before=0, line=None, first=False, upper=False,
         spacing=None):
    """runs: str or list of (text, color, bold, font) tuples."""
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment = align
    _set_space_after(p, space_after)
    if space_before:
        p.space_before = Pt(space_before)
    if line:
        p.line_spacing = line
    if isinstance(runs, str):
        runs = [(runs, color, bold, font)]
    for item in runs:
        text, c, b, f = (item + (None, None, None))[:4]
        if upper:
            text = text.upper()
        r = p.add_run()
        r.text = text
        r.font.size = Pt(size)
        r.font.name = f or font
        r.font.bold = bool(b) if b is not None else bold
        r.font.color.rgb = c if c is not None else (color or INK)
        if spacing is not None:
            rPr = r._r.get_or_add_rPr()
            rPr.set('spc', str(int(spacing)))  # letter spacing in 1/100 pt
    return p


def brandbar(slide, section, dark=False):
    lock = MUTED if not dark else ONINKMUT
    name_c = INK if not dark else ONINK
    tf = tx(slide, ML, 0.5, CW * 0.6, 0.4)
    para(tf, [("Digital", name_c, True, DISPLAY), (" Experiments", lock, False, DISPLAY)],
         13, first=True, space_after=0)
    tf2 = tx(slide, ML + CW * 0.6, 0.5, CW * 0.4, 0.4)
    para(tf2, section, 10.5, lock, font=MONO, align=PP_ALIGN.RIGHT, first=True,
         space_after=0, upper=False, spacing=120)


def pagenum(slide, n, dark=False):
    tf = tx(slide, SW - ML - 1.2, SH - 0.72, 1.2, 0.35)
    para(tf, f"{n:02d}", 10, MUTED2 if not dark else ONINKMUT, font=MONO,
         align=PP_ALIGN.RIGHT, first=True, space_after=0)


def eyebrow(tf, text, dark=False, first=False):
    para(tf, text.upper(), 11.5, ACCENT3 if dark else ACCENT, font=MONO,
         first=first, space_after=10, spacing=180)


def card(slide, l, t, w, h, fill, line_c):
    sp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(l), Inches(t),
                                Inches(w), Inches(h))
    sp.fill.solid(); sp.fill.fore_color.rgb = fill
    sp.line.color.rgb = line_c; sp.line.width = Pt(1)
    sp.shadow.inherit = False
    # softer corner radius
    try:
        sp.adjustments[0] = 0.06
    except Exception:
        pass
    tf = sp.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.28)
    tf.margin_top = tf.margin_bottom = Inches(0.26)
    return tf


def bullets(tf, items, dark=False, size=14, first_blank=True):
    body = ONINKMUT if dark else MUTED
    for i, it in enumerate(items):
        para(tf, [("•  ", ACCENT3 if dark else ACCENT, False, SANS),
                  (it, body, False, SANS)],
             size, first=(i == 0 and not first_blank), space_after=5, line=1.05)


# ============================================================ SLIDES =========

# 01 · COVER -----------------------------------------------------------------
s = slide(INK); brandbar(s, "Capabilities", dark=True)
tf = tx(s, ML, 2.0, CW, 3.6)
eyebrow(tf, "Bespoke Game UX & Player Research", dark=True, first=True)
para(tf, [("better design,", ONINK, False, DISPLAY)], 54, line=1.02, space_after=0)
para(tf, [("better games.", ONINK, False, DISPLAY)], 54, line=1.02, space_after=14)
para(tf, "Digital Experiments is a third-party developer providing bespoke user "
         "experience design for video game studios — from tactical UX/UI execution "
         "to senior-level strategic UX.", 18, ONINKMUT, line=1.3, space_after=14)
para(tf, "www.digitalexperiments.com", 14, ONINKMUT, font=MONO)
pagenum(s, 1, dark=True)

# 02 · CAPABILITIES ----------------------------------------------------------
s = slide(PAPER); brandbar(s, "Capabilities")
tf = tx(s, ML, 1.35, CW, 1.4)
eyebrow(tf, "Capabilities", first=True)
para(tf, "Everything from tactical execution to strategic UX.", 30, INK,
     font=DISPLAY, line=1.05)
cols = [
    ("Tactical UI/UX", ["High Fidelity Prototyping", "Wireframing / Greyboxing",
                        "UX Design", "UI Design", "UI Animation", "UI Implementation"]),
    ("Strategic UX", ["Vision Clarification", "Vision Deck Design", "Game Brand Design",
                      "Product / Market Fit", "Persona Development", "Design Systems"]),
    ("UX for Teams", ["LEAN UCD Training", "Maturing UX Practices", "Psychological Safety",
                      "Collaboration Workshops", "Process Design", "Service Design"]),
]
cw, gap = 3.7, 0.42
top, ch = 3.05, 3.5
for i, (title, items) in enumerate(cols):
    l = ML + i * (cw + gap)
    ctf = card(s, l, top, cw, ch, PAPER, LINE)
    para(ctf, title.upper(), 12.5, ACCENT, font=MONO, first=True, space_after=12, spacing=80)
    bullets(ctf, items, size=14)
pagenum(s, 2)

# 03 · CLIENTS & GAMES -------------------------------------------------------
s = slide(PAPER2); brandbar(s, "Clients & games")
tf = tx(s, ML, 1.5, CW, 1.6)
eyebrow(tf, "Clients & games", first=True)
para(tf, "Trusted on shipped and unreleased titles across the industry.", 30, INK,
     font=DISPLAY, line=1.05)
clients = ["Activision Blizzard", "SciPlay", "Mythical Games", "38 Studios"]
cw, gap, top, ch = 2.74, 0.34, 3.2, 1.5
for i, name in enumerate(clients):
    l = ML + i * (cw + gap)
    ctf = card(s, l, top, cw, ch, PAPER, LINE)
    ctf.vertical_anchor = MSO_ANCHOR.MIDDLE
    para(ctf, name, 16, INK2, font=DISPLAY, align=PP_ALIGN.CENTER, first=True, space_after=0)
tf2 = tx(s, ML, 5.05, CW, 1.2)
para(tf2, "Engagements span AAA, social casino, and MMO — including work on titles such "
          "as Jackpot Party (SciPlay) and Blankos Block Party (Mythical Games). "
          "NDA-ready and experienced with unreleased IP.", 13, MUTED, line=1.3, first=True)
pagenum(s, 3)

# 04 · OUR TEAM --------------------------------------------------------------
s = slide(INK); brandbar(s, "Our team", dark=True)
tf = tx(s, ML, 1.35, CW, 2.2)
eyebrow(tf, "Our team", dark=True, first=True)
para(tf, "Digital Experiments is a third party developer that provides bespoke user "
         "experience design for video game studios. Staffed by a veteran team with over "
         "20 years of UX experience in games, DXP offers everything from tactical UX/UI "
         "execution as well as more senior level strategic UX work.",
     20, ONINK, font=DISPLAY, line=1.25)
roster = [
    ("Irena Pereira", "Director · Lead Designer / Engineer"),
    ("SJ Dalmar", "UI/UX Designer · Animator"),
    ("Alyssa Yeo", "UI/UX Designer"),
    ("Kelsey Phelan", "UI/UX Designer · Research"),
    ("Hannah Kinzinger", "Engineer · Gameplay"),
    ("Leslie Crystal", "Associate Producer · Operations"),
]
cw, gap, ch = 3.7, 0.42, 0.95
top0 = 4.25
for i, (name, role) in enumerate(roster):
    r, c = divmod(i, 3)
    l = ML + c * (cw + gap)
    t = top0 + r * (ch + 0.28)
    ctf = card(s, l, t, cw, ch, CARDINK, ONINKLN)
    ctf.vertical_anchor = MSO_ANCHOR.MIDDLE
    para(ctf, name, 16, ONINK, font=DISPLAY, bold=True, first=True, space_after=3)
    para(ctf, role.upper(), 9.5, ONINKMUT, font=MONO, spacing=60)
pagenum(s, 4, dark=True)

# 05 · TOOLS -----------------------------------------------------------------
s = slide(PAPER); brandbar(s, "Design team · tools")
tf = tx(s, ML, 2.2, CW, 2.0)
eyebrow(tf, "Design team · tools", first=True)
para(tf, "A senior, cross-discipline team — designers, engineers & producers.",
     30, INK, font=DISPLAY, line=1.08, space_after=18)
tools = ["Figma", "Adobe Creative Suite", "Unity", "C# & multiple languages",
         "Visual Studio", "Google & Office Suite"]
# chips row as one paragraph of pill-ish text
ctf = tx(s, ML, 4.1, CW, 1.0)
para(ctf, "   ".join(tools), 16, INK2, font=SANS, first=True, line=1.6)
pagenum(s, 5)


# ---- Bio slides ------------------------------------------------------------
def bio_slide(n, idx, name, role, body, skills, hobbies, faves, dark_soft):
    bg = PAPER2 if dark_soft else PAPER
    s = slide(bg); brandbar(s, f"Team · {idx} / 6")
    hd = tx(s, ML, 1.25, CW, 1.3)
    para(hd, name, 34, INK, font=DISPLAY, first=True, space_after=4)
    para(hd, role.upper(), 12, ACCENT, font=MONO, spacing=80)
    bt = tx(s, ML, 2.65, CW, 2.1)
    para(bt, body, 15, INK2, line=1.5, first=True)
    # mini columns
    cols = [("Skills", skills), ("Hobbies", hobbies), ("Favorite Games", faves)]
    cw, gap, top = 3.7, 0.42, 5.1
    for i, (h, items) in enumerate(cols):
        l = ML + i * (cw + gap)
        mt = tx(s, l, top, cw, 2.0)
        para(mt, h.upper(), 11, ACCENT, font=MONO, first=True, space_after=8, spacing=120)
        for it in items:
            para(mt, it, 12, MUTED, line=1.05, space_after=4)
    pagenum(s, n)


bio_slide(6, 1, "Irena Pereira",
          "Director · Product / Design / Engineering · Product Leader",
          "From World of Warcraft to REALTOR.com®, Irena has spent an over 20-year "
          "career as an engineer, NATSEC product designer, and game designer. Irena has "
          "been a trailblazer in UX/UI where her work was instrumental in defining the "
          "field in our industry. She pioneered work methodologies we take for granted "
          "today, such as A/B testing and data-driven design. She is a speaker, designer, "
          "and educator. Today she helps game studios and B2C organizations define and "
          "craft groundbreaking user experiences to increase engagement and revenue.",
          ["Game Design", "Strategic UX", "Team Leadership", "Alignment", "Prioritization"],
          ["Photography", "Mycology", "Hiking", "Diplomacy (Board Game)"],
          ["Final Fantasy VII", "EverQuest", "World of Warcraft", "Diablo", "Civilization II"],
          False)

bio_slide(7, 2, "SJ Dalmar", "UI/UX Designer · Animation / 2D Illustration",
          "SJ has been fascinated by art their whole life, especially animation. After "
          "graduating from VanArts in 2016, they worked as a character animator at Bardel "
          "Entertainment before stepping into UX/UI design. Their animation work combined "
          "with an eye for composition honed by 15+ years of illustration gives them a "
          "unique design perspective. SJ accomplished creating a task management app "
          "focusing on alleviating challenges for folks with ADHD. Today, as a designer "
          "with Digital Experiments, they’re focused on designing delightful and "
          "accessible play experiences from their home in the Pacific Northwest.",
          ["Wireframing", "Prototyping", "Data Visualization", "Motion Graphics", "Character Design"],
          ["Taxidermy", "TTRPGs", "Sewing"],
          ["Persona 4", "Dragon Age", "Sekiro", "Genshin Impact"],
          True)

bio_slide(8, 3, "Alyssa Yeo", "UI/UX Designer · Game Development / Tech",
          "Alyssa first started out her product design journey in corporate CX, having "
          "co-facilitated and strategized workshops for the banking sector in Kuala Lumpur, "
          "Malaysia. She’s been sketching, designing, iterating and reiterating "
          "prototypes for most of her professional life, drawing from her background in "
          "human psychology and natural instincts for branding. Her pivot into video games "
          "began with a pandemic-induced dabble into TTRPG design where she reinvigorated "
          "her love for narrative writing and producing. With Digital Experiments, Alyssa "
          "has completed indie UX projects for startups and gamedev companies, notably "
          "Mythical Games and SciPlay. On the side, Alyssa is exploring UI implementation, "
          "while indulging in game writing with a love for making unlikely combinations of "
          "seemingly incompatible genres.",
          ["UI / UX Design", "Narrative Design", "High–Low Fidelity Prototyping",
           "Adobe Creative Suite, Figma"],
          ["Designing tabletop story games", "Antique / vintage fashions", "Theatre",
           "Big meals from a tiny kitchen"],
          ["Red Dead Redemption 2", "Disco Elysium", "Inside", "Dragon Age"],
          False)

bio_slide(9, 4, "Kelsey Phelan", "UI/UX Designer · Research / Game Development",
          "Kelsey has always had a fascination with research. During her Biomedical "
          "Engineering master’s program, she conducted research with Alzheimer’s "
          "Disease patients. It became clear that a person’s interactions with "
          "technology would be a limiting factor in their understanding. From there, Kelsey "
          "delved into UX Design to better understand how we could utilize technology as a "
          "society to create a more understanding and inclusive culture overall. Her "
          "background in research has provided valuable insight for finding research "
          "artifacts, a go-getter attitude and crafting solutions. Her UX work has focused "
          "on XR apps for science study, a social scavenger hunt game, and other games.",
          ["Wireframing", "UX Design", "Data Visualization", "Research Synthesis"],
          ["Activity map of the city", "Aerial Arts", "DM of 5E DnD", "Animations"],
          ["Persona 3", "Chrono Trigger", "Shadow of the Colossus", "Psychonauts"],
          True)

bio_slide(10, 5, "Hannah Kinzinger", "Engineer · Gameplay / Generalist",
          "A talented software engineer with a broad range of knowledge in multiple "
          "disciplines, Hannah has worn several hats in her game design career. She lead an "
          "ad hoc indie game team on designing and developing several projects using "
          "flexible, scalable software solutions in a high pressure environment. Her work "
          "teaching programming fundamentals to hundreds of students between the ages of 11 "
          "and 14 honed her communication skills. Hannah brings all of this and more to "
          "Digital Experiments, where she has coordinated, implemented, and tested software "
          "systems on upcoming titles.",
          ["Unity", "C# and multiple other languages", "Flexible Interdisciplinary "
           "Communicator", "Visual Studio"],
          ["Tabletop RPGs", "Digital Painting", "Stargazing"],
          ["Final Fantasy XIV", "Elden Ring", "Monster Hunter", "Terra Nil"],
          False)

bio_slide(11, 6, "Leslie Crystal", "Associate Producer · Operations Manager",
          "With a lifelong love for RPG’s and nearly two decades as a paralegal + "
          "legal ghostwriter (intellectual property law, employment law, corporate "
          "compliance/business operations, et al.), Leslie pivoted into the world of game "
          "design and development in 2020. Since then, Leslie has served as the production "
          "scrum master for Tochi (Snackpass), B4CKP4CK (Studio Rebase), Begone Beast "
          "(Studio Tandemi), Town of Zoz (Studio Pixanoh), and a number of unannounced "
          "projects under game designer Alexander Brazie and Atomech LLC. Leslie currently "
          "serves as the Operations Manager for Digital Experiments, providing the "
          "organizational support and administrative framework to a talented group of "
          "artists, engineers, and designers headed by Director and UX/UI Trailblazer, "
          "Irena Pereira.",
          ["Team Organization", "Project Production & Development",
           "Legal / Business Operations", "Google, Adobe & Office Suite"],
          ["Tarot & The Arcane", "Homesteading", "Handcrafted Cosmetics", "Embroidery"],
          ["World of Warcraft", "Witchy Life Story", "The Witcher", "Wylde Flowers"],
          True)

# 12 · CLOSING ---------------------------------------------------------------
s = slide(ACCENT); brandbar(s, "Let’s talk", dark=True)
tf = tx(s, ML, 2.4, CW, 3.0)
para(tf, "LET’S BUILD SOMETHING PLAYERS LOVE", 11.5, RGBColor(0xFF, 0xFF, 0xFF),
     font=MONO, first=True, space_after=12, spacing=180)
para(tf, "better design, better games.", 50, RGBColor(0xFF, 0xFF, 0xFF),
     font=DISPLAY, line=1.05, space_after=16)
para(tf, "sayhi@digitalexperiments.com    ·    www.digitalexperiments.com", 15,
     RGBColor(0xFF, 0xFF, 0xFF), font=MONO, space_after=10)
para(tf, "Digital Experiments — bespoke Game UX & Player Research", 12,
     RGBColor(0xEC, 0xEA, 0xFB), font=SANS)
pagenum(s, 12, dark=True)

out = os.path.join(os.path.dirname(__file__), "Digital-Experiments-Capabilities.pptx")
prs.save(out)
print("Saved", out, "·", len(prs.slides.__iter__.__self__._sldIdLst), "slides")
